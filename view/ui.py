from datetime import datetime
from io import BytesIO
import os
import platform
import subprocess
from tkinter import messagebox

import pandas as pd
import ttkbootstrap as ttk
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill

from src.main import main
from src.memory import MemoryManager
from src.metrics import Metrics
from src.schedulers import CPU_Scheduler
from src.services.data_loader import DataLoader
from view.components import DataFrameTable, GanttView, SimulationControls
from view.metrics_view import MetricsView
from view.output_format import (
    DEVICE_COLORS,
    format_device_output,
    format_process_names,
    process_color_map,
)


APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PROCESS_FILE = os.path.join(APP_DIR, "data", "cpu", "processes.json")
PAGES_FILE = os.path.join(APP_DIR, "data", "memory", "memory_requering.json")
ACL_FILE = os.path.join(APP_DIR, "data", "access_control", "access_control_list.json")
DEVICES_FILE = os.path.join(APP_DIR, "data", "devices", "devices_list.json")
DOCS_DIR = os.path.join(APP_DIR, "docs")
CPU_METHODS = CPU_Scheduler.Methods
MEMORY_METHODS = MemoryManager.Methods


class SimulatorApp:
    # Prepara a janela, os dados fixos e os componentes da interface.
    def __init__(self, window):
        self.window = window
        self.window.title("KernelLab")
        self.window.geometry("1350x900")

        self.last_results = None
        self.last_settings = None
        self.processes = DataLoader.read_dataframe(PROCESS_FILE)
        process_names = []
        for process_name in self.processes["name"]:
            process_names.append(process_name)
        self.process_names = []
        for name in process_names:
            self.process_names.append(name.upper())
        self.process_colors = process_color_map(process_names)
        self._build_window()

    # Cria os controles, as abas e os botões da janela.
    def _build_window(self):
        cpu_method_names = []
        for method in CPU_METHODS:
            cpu_method_names.append(method.name)

        memory_method_names = []
        for method in MEMORY_METHODS:
            memory_method_names.append(method.name)

        self.controls = SimulationControls(
            self.window, cpu_method_names, memory_method_names
        )

        buttons = ttk.Frame(self.window, padding=(8, 0, 8, 8))
        buttons.pack(anchor="center")
        ttk.Button(
            buttons, text="Executar", command=self.run_simulation, bootstyle="success"
        ).pack(side="left", padx=(0, 5))
        self.save_button = ttk.Button(
            buttons,
            text="Salvar resultados",
            command=self.save_results,
            state="disabled",
        )
        self.save_button.pack(side="left", padx=5)
        ttk.Button(
            buttons, text="Abrir pasta docs", command=self.open_docs_folder
        ).pack(side="left", padx=5)

        self.notebook = ttk.Notebook(self.window)
        self.notebook.pack(fill="both", expand=True, padx=8, pady=4)

        self.metrics_view = MetricsView(self.notebook, self.process_colors)

        self.cpu_schedule_table = DataFrameTable(
            self.notebook,
            "Escalonamento CPU",
            "process",
            row_colors=self.process_colors,
        )
        self.memory_schedule_table = DataFrameTable(self.notebook, "Estado da memória")
        self.devices_table = DataFrameTable(
            self.notebook, "Dispositivos", "device", row_colors=DEVICE_COLORS
        )
        self.gantt_view = GanttView(self.notebook, self.process_colors)

    # Lê a sequência padrão ou monta a sequência digitada pelo usuário.
    def _get_pages(self) -> pd.DataFrame:
        if self.controls.pages_mode.get() == "default":
            return DataLoader.read_dataframe(PAGES_FILE)

        pages = []
        text = self.controls.pages_entry.get()
        for page in text.split(","):
            page = page.strip()
            if page != "":
                pages.append(page)

        if len(pages) == 0:
            raise ValueError("Digite as páginas separadas por vírgula.")
        return pd.DataFrame({"page": pages})

    # Lê as escolhas da interface e executa a simulação.
    def run_simulation(self):
        try:
            cpu_method = CPU_METHODS[self.controls.cpu_choice.get()]
            memory_method = MEMORY_METHODS[self.controls.memory_choice.get()]
            quantum = int(self.controls.quantum.get())
            frames = int(self.controls.frames.get())
            if quantum <= 0:
                raise ValueError("O quantum deve ser maior que zero.")
            if frames <= 0:
                raise ValueError("A quantidade de frames deve ser maior que zero.")
            pages = self._get_pages()

            cpu_schedule, memory_schedule, device_schedule = main(
                PROCESS_FILE,
                pages,
                ACL_FILE,
                DEVICES_FILE,
                cpu_method,
                memory_method,
                memory_capacity=frames,
                quantum=quantum,
            )

            page_sequence = []
            for page in pages["page"]:
                page_sequence.append(str(page))
            settings = {
                "cpu_method": cpu_method.name,
                "memory_method": memory_method.name,
                "quantum": quantum,
                "frames": frames,
                "page_sequence": ", ".join(page_sequence),
            }
            self._show_results(
                cpu_schedule,
                memory_schedule,
                device_schedule,
                settings,
            )
        except Exception as error:
            messagebox.showerror("Erro", str(error), parent=self.window)

    # Calcula, guarda e exibe os resultados da simulação.
    def _show_results(
        self,
        cpu_schedule,
        memory_schedule,
        device_schedule,
        settings,
    ):
        cpu_metrics = format_process_names(Metrics.cpu(self.processes, cpu_schedule))
        memory_metrics = Metrics.memory(memory_schedule)
        gantt_data = format_process_names(Metrics.gantt(cpu_schedule))
        device_output = format_device_output(device_schedule)

        self.last_results = {
            "CPU_Metricas": cpu_metrics,
            "CPU_Escalonamento": gantt_data,
            "Memoria_Metricas": memory_metrics,
            "Memoria_Estado": memory_schedule,
            "Dispositivos": device_output,
            "Gantt_CPU": gantt_data,
        }
        self.last_settings = pd.DataFrame([settings])

        self.metrics_view.show(cpu_metrics, memory_metrics)
        self.cpu_schedule_table.show(gantt_data)
        self.memory_schedule_table.show(memory_schedule)
        self.devices_table.show(device_output)
        self.gantt_view.show(gantt_data, self.process_names)
        self.notebook.select(self.metrics_view.cpu_table.table.winfo_parent())
        self.save_button.configure(state="normal")

    # Salva as tabelas e o gráfico em uma planilha Excel.
    def save_results(self):
        if self.last_results is None:
            return

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        file_path = os.path.join(DOCS_DIR, f"resultados_{timestamp}.xlsx")
        image_data = BytesIO()

        try:
            os.makedirs(DOCS_DIR, exist_ok=True)
            self.gantt_view.fig.savefig(image_data, format="png")
            image_data.seek(0)

            with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
                self._write_workbook(writer, image_data)

            messagebox.showinfo(
                "Salvo", f"Arquivo salvo em:\n{file_path}", parent=self.window
            )
        except Exception as error:
            messagebox.showerror("Erro ao salvar", str(error), parent=self.window)
        finally:
            image_data.close()

    # Escreve as tabelas e aplica a formatação das planilhas.
    def _write_workbook(self, writer, image_data):
        self.last_settings.to_excel(writer, sheet_name="Configuracao", index=False)
        for sheet_name, dataframe in self.last_results.items():
            dataframe.to_excel(writer, sheet_name=sheet_name, index=False)
        writer.sheets["Gantt_CPU"].add_image(ExcelImage(image_data), "E2")
        self._format_workbook(writer)

    # Padroniza cabeçalhos, alinhamento e cores das planilhas.
    def _format_workbook(self, writer):
        header_fill = PatternFill(fill_type="solid", fgColor="E7E6E6")
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            for row in worksheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(
                        horizontal="center", vertical="center"
                    )
            for cell in worksheet[1]:
                if cell.value is not None:
                    cell.value = str(cell.value).upper()
                cell.fill = header_fill
                cell.font = Font(bold=True)

        self._color_workbook_rows(writer)

    # Colore as linhas de processos e dispositivos nas planilhas.
    def _color_workbook_rows(self, writer):
        process_style = {
            "name_column": "PROCESS",
            "colors": self.process_colors,
            "bold_columns": ["PROCESS"],
        }
        device_style = {
            "name_column": "DEVICE",
            "colors": DEVICE_COLORS,
            "bold_columns": ["DEVICE", "RUNNING_PROCESS", "QUEUE"],
        }
        style_sheets = {
            "CPU_Metricas": process_style,
            "CPU_Escalonamento": process_style,
            "Gantt_CPU": process_style,
            "Dispositivos": device_style,
        }
        for sheet_name, sheet_style in style_sheets.items():
            self._color_sheet_rows(writer.sheets[sheet_name], sheet_style)

    # Aplica as cores de uma tabela conforme o processo ou dispositivo.
    def _color_sheet_rows(self, worksheet, style):
        name_column = style["name_column"]
        colors = style["colors"]
        bold_columns = style["bold_columns"]
        name_columns = {}
        for cell in worksheet[1]:
            name_columns[cell.value] = cell.column

        entity_column = name_columns[name_column]
        for row_number in range(2, worksheet.max_row + 1):
            entity_name = worksheet.cell(
                row=row_number, column=entity_column
            ).value
            if entity_name not in colors:
                continue

            color = colors[entity_name]
            row_fill = PatternFill(
                fill_type="solid", fgColor=color.lstrip("#")
            )
            for cell in worksheet[row_number]:
                cell.fill = row_fill
            for column_name in bold_columns:
                column_number = name_columns[column_name]
                worksheet.cell(
                    row=row_number, column=column_number
                ).font = Font(bold=True)

    # Abre a pasta onde os arquivos Excel são salvos.
    def open_docs_folder(self):
        try:
            os.makedirs(DOCS_DIR, exist_ok=True)
            if platform.system() == "Windows":
                os.startfile(str(DOCS_DIR))
            elif platform.system() == "Darwin":
                subprocess.Popen(["open", str(DOCS_DIR)])
            else:
                subprocess.Popen(["xdg-open", str(DOCS_DIR)])
        except Exception as error:
            messagebox.showerror("Erro ao abrir pasta", str(error), parent=self.window)


# Cria e inicia a janela principal do simulador.
def run_app():
    window = ttk.Window(themename="darkly")
    SimulatorApp(window)
    window.mainloop()
